from openpyxl import load_workbook
import os

path = r"d:\Proyecto2026\ACP_DWH\ACP Proyecciones\ETL\data\entrada\evaluacion_pesos\Pesos 2023.xlsx"
wb = load_workbook(filename=path, read_only=True, data_only=True)
sheet = wb.active
for row in sheet.iter_rows(values_only=True):
    print("Columnas 2023:", row)
    break
wb.close()
