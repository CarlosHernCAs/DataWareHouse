import os
from openpyxl import load_workbook

def count_rows(file_path):
    try:
        wb = load_workbook(filename=file_path, read_only=True, data_only=True)
        sheet = wb.active
        count = 0
        for row in sheet.iter_rows(values_only=True):
            if any(row):
                count += 1
        wb.close()
        return count
    except Exception as e:
        return f"Error: {e}"

base_dir = r"d:\Proyecto2026\ACP_DWH\ACP Proyecciones\ETL\data\entrada\evaluacion_pesos"
files = [f for f in os.listdir(base_dir) if f.endswith('.xlsx')]

for f in files:
    path = os.path.join(base_dir, f)
    rows = count_rows(path)
    print(f"{f}: {rows} filas reales")
