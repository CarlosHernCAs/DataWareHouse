import os
from openpyxl import load_workbook

def read_excel_fast(file_path):
    wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    sheet = wb.active
    data = []
    headers = []
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0:
            headers = list(row)
        else:
            if not any(row): continue
            data.append(list(row))
    wb.close()
    return headers, data

base_dir = r"C:\Users\chernandez\Desktop\historico"
h_2021, d_2021 = read_excel_fast(os.path.join(base_dir, "Pesos 2021.xlsx"))
h_2022, d_2022 = read_excel_fast(os.path.join(base_dir, "Pesos 2022.xlsx"))

# Calcular la suma de Peso Total en cada archivo
idx_peso = h_2021.index('Peso Total ')

def sum_peso(data):
    total = 0
    for r in data:
        val = r[idx_peso]
        if isinstance(val, (int, float)):
            total += val
    return total

suma_2021 = sum_peso(d_2021)
suma_2022 = sum_peso(d_2022)

print(f"Suma de 'Peso Total' en 2021: {suma_2021}")
print(f"Suma de 'Peso Total' en 2022: {suma_2022}")
print(f"Suma de 'Peso Total' si simplemente unimos todo: {suma_2021 + suma_2022}")

# Analisis de duplicados
t_2021 = [tuple(str(x) for x in r) for r in d_2021]
t_2022 = [tuple(str(x) for x in r) for r in d_2022]

set_2021 = set(t_2021)
set_2022 = set(t_2022)

interseccion = set_2021.intersection(set_2022)
print(f"Total filas 2021: {len(t_2021)}")
print(f"Total filas 2022: {len(t_2022)}")
print(f"Filas de 2022 que ya existen exactamente en 2021: {len(interseccion)}")

print("\nEjemplo de 3 filas duplicadas:")
for i, t in enumerate(list(interseccion)[:3]):
    print(f"Dup {i+1}: Semana {t[h_2021.index('SEMANA')]}, Modulo {t[h_2021.index('M')]}, Variedad {t[h_2021.index('VARIEDAD')]}, Peso {t[idx_peso]}")
