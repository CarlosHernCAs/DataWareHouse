import pandas as pd
import os
import sys
from openpyxl import load_workbook
import datetime

log_file = r"d:\Proyecto2026\ACP_DWH\ACP Proyecciones\scratch\log_pesos.txt"
def p(msg):
    with open(log_file, "a") as f:
        f.write(f"[{datetime.datetime.now()}] {msg}\n")
    print(msg, flush=True)

open(log_file, "w").close() # Clear log

def read_excel_fast(file_path):
    p(f"Abriendo {file_path} con openpyxl...")
    wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    sheet = wb.active
    data = []
    headers = []
    
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0:
            headers = list(row)
        else:
            # Si la fila entera es None o vacia, ignorarla
            if not any(row):
                continue
            data.append(list(row))
            
    wb.close()
    return pd.DataFrame(data, columns=headers)

try:
    p("Iniciando script...")
    base_dir = r"C:\Users\chernandez\Desktop\historico"
    file_2021 = os.path.join(base_dir, "Pesos 2021.xlsx")
    file_2022 = os.path.join(base_dir, "Pesos 2022.xlsx")
    output_file = os.path.join(base_dir, "Pesos_Consolidados_2021_2022.xlsx")

    df_2021 = read_excel_fast(file_2021)
    p(f"2021 leido. Filas reales: {len(df_2021)}")

    df_2022 = read_excel_fast(file_2022)
    p(f"2022 leido. Filas reales: {len(df_2022)}")

    p("Combinando dataframes...")
    df_concat = pd.concat([df_2021, df_2022], ignore_index=True)

    filas_antes = len(df_concat)
    p(f"Eliminando duplicados de {filas_antes} filas...")
    # Asegurar tipos comparables
    df_concat = df_concat.astype(str)
    df_concat.drop_duplicates(inplace=True)
    filas_despues = len(df_concat)
    p(f"Filas despues de duplicados: {filas_despues}")

    p("Mapeando columnas...")
    df_final = pd.DataFrame()
    df_final['Semana'] = df_concat['SEMANA']
    df_final['Campaña'] = df_concat['CAMPAÑA']
    df_final['Fecha'] = None
    df_final['DNI'] = None
    df_final['Nombres'] = None
    df_final['M'] = df_concat['M']
    df_final['T'] = df_concat['T']
    df_final['Valvula'] = None
    df_final['VAR.'] = df_concat['VARIEDAD']
    df_final['Cosechables'] = df_concat['Bayas']
    df_final['Peso Cosechables'] = df_concat['Peso Total ']
    df_final['Etapa Cosecha'] = None
    df_final['AÑO'] = df_concat['AÑO']
    df_final['Mes'] = None

    p("Guardando a Excel...")
    df_final.to_excel(output_file, index=False)
    p(f"Exito. Guardado en {output_file}")
except Exception as e:
    p(f"ERROR FATAL: {e}")
