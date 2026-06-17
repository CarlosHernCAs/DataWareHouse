import sys
from openpyxl import load_workbook

def print_cols(file_path):
    print(f"\n--- Columns in {file_path} ---")
    try:
        wb = load_workbook(filename=file_path, read_only=True, data_only=True)
        sheet = wb.active
        
        # Get first two rows
        rows = sheet.iter_rows(min_row=1, max_row=2, values_only=True)
        headers = next(rows, [])
        sample_data = next(rows, [])
        
        print(f"Headers: {headers}")
        print(f"Sample Data: {sample_data}")
        wb.close()
    except Exception as e:
        print(f"Error reading {file_path}: {e}")

base_dir = r"C:\Users\chernandez\Desktop\historico"
print_cols(base_dir + r"\Pesos 2021.xlsx")
print_cols(base_dir + r"\Pesos 2022.xlsx")
print_cols(base_dir + r"\Pesos 2023.xlsx")
