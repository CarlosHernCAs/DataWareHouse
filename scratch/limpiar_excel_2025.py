import os
import pandas as pd
from openpyxl import load_workbook

def read_excel_fast(file_path):
    wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    sheet = wb.active
    data = []
    headers = []
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0:
            headers = list(row)
        else:
            if not any(row):
                continue
            data.append(list(row))
    wb.close()
    return pd.DataFrame(data, columns=headers)

path = r"d:\Proyecto2026\ACP_DWH\ACP Proyecciones\ETL\data\entrada\evaluacion_pesos\Evaluación de pesos_2025.xlsx"
print(f"Limpiando archivo gigante: {path}")
df = read_excel_fast(path)
print(f"Leidas {len(df)} filas reales.")
df.to_excel(path, index=False)
print("Archivo sobreescrito y limpio.")
