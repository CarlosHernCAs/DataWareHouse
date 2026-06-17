import os
import sys
from openpyxl import load_workbook, Workbook
import datetime

log_file = r"d:\Proyecto2026\ACP_DWH\ACP Proyecciones\scratch\log_pesos.txt"
def p(msg):
    with open(log_file, "a") as f:
        f.write(f"[{datetime.datetime.now()}] {msg}\n")
    print(msg, flush=True)

open(log_file, "w").close()

def read_excel_fast(file_path):
    p(f"Abriendo {file_path} con openpyxl...")
    wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    sheet = wb.active
    data = []
    headers = []
    
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0:
            headers = list(row)
        else:
            if not any(row):
                continue
            data.append(list(row))
            
    wb.close()
    return headers, data

try:
    p("Iniciando script openpyxl puro...")
    base_dir = r"C:\Users\chernandez\Desktop\historico"
    file_2021 = os.path.join(base_dir, "Pesos 2021.xlsx")
    file_2022 = os.path.join(base_dir, "Pesos 2022.xlsx")
    output_file = os.path.join(base_dir, "Pesos_Consolidados_2021_2022.xlsx")

    h_2021, d_2021 = read_excel_fast(file_2021)
    p(f"2021 leido. Filas reales: {len(d_2021)}")

    h_2022, d_2022 = read_excel_fast(file_2022)
    p(f"2022 leido. Filas reales: {len(d_2022)}")

    p("Combinando listas y eliminando duplicados exactos...")
    todas_las_filas = d_2021 + d_2022
    
    # Eliminar duplicados usando set (convertir filas a tuplas de string para comparar)
    filas_unicas = []
    vistos = set()
    for row in todas_las_filas:
        t_row = tuple(str(x) for x in row)
        if t_row not in vistos:
            vistos.add(t_row)
            filas_unicas.append(row)
            
    p(f"Total combinadas: {len(todas_las_filas)}. Unicas: {len(filas_unicas)}")

    # Headers del nuevo archivo
    # 'Semana', 'Campaña', 'Fecha', 'DNI', 'Nombres', 'M', 'T', 'Valvula', 'VAR.', 
    # 'Cosechables', 'Peso Cosechables', 'Etapa Cosecha', 'AÑO', 'Mes'
    new_headers = ['Semana', 'Campaña', 'Fecha', 'DNI', 'Nombres', 'M', 'T', 'Valvula', 'VAR.', 'Cosechables', 'Peso Cosechables', 'Etapa Cosecha', 'AÑO', 'Mes']
    
    # Crear nuevo workbook
    out_wb = Workbook()
    out_sheet = out_wb.active
    out_sheet.append(new_headers)

    # Mapeo:
    # 2021 Headers: 'CAMPAÑA', 'AÑO', 'SEMANA', 'M', 'T', 'VARIEDAD', '#ORGANOS', 'Bayas', 'Peso Total ', 'PESO', 'Peso 2', 'KILOS COSECHADOS'
    # indices en 2021/2022:
    idx_semana = h_2021.index('SEMANA')
    idx_campana = h_2021.index('CAMPAÑA')
    idx_ano = h_2021.index('AÑO')
    idx_m = h_2021.index('M')
    idx_t = h_2021.index('T')
    idx_var = h_2021.index('VARIEDAD')
    idx_bayas = h_2021.index('Bayas')
    idx_peso = h_2021.index('Peso Total ')

    for row in filas_unicas:
        new_row = [
            row[idx_semana],   # Semana
            row[idx_campana],  # Campaña
            None,              # Fecha
            None,              # DNI
            None,              # Nombres
            row[idx_m],        # M
            row[idx_t],        # T
            None,              # Valvula
            row[idx_var],      # VAR.
            row[idx_bayas],    # Cosechables
            row[idx_peso],     # Peso Cosechables
            None,              # Etapa Cosecha
            row[idx_ano],      # AÑO
            None               # Mes
        ]
        out_sheet.append(new_row)

    p("Guardando a Excel...")
    out_wb.save(output_file)
    p(f"Exito. Guardado en {output_file}")
except Exception as e:
    p(f"ERROR FATAL: {e}")
